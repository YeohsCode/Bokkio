import asyncio
import copy
import hashlib
import importlib.metadata
import json
from pathlib import Path
import subprocess
import sys

import pytest

from bokkio import capability_registry as registry
from bokkio.runtime import Runtime
from bokkio.sdk import Client


@pytest.fixture
def setup(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    workbook.active.append(["Amount", "Owner"])
    workbook.active.append([320, "林安"])
    workbook.save(path)
    workbook.close()
    policy = {"paths": [str(tmp_path)], "modules": [registry.MODULE],
              "write": True, "trusted_host": True, "require_approval": False}
    client = Client(Runtime(tmp_path / "state", policy))
    client.open()
    yield client, path, policy
    client.call("session.close")


def inputs(path, **kwargs):
    return {"path": str(path), "sheet": "Summary", "expected_sha256": hashlib.sha256(path.read_bytes()).hexdigest(), **kwargs}


def request(client, cid, data):
    contract = client.call("capability.get", {"capability_id": cid})["result"]
    return {"capability_id": cid, "version": contract["version"], "contract_sha256": contract["contract_sha256"], "inputs": data}


def invoke(client, req, op=None):
    queued = client.call("capability.invoke", req, op)
    assert queued["status"] == "queued", queued
    final = client.wait_job(queued["result"]["job_id"], timeout=20)
    return queued, final["result"]


def test_discovery_is_progressive_and_default_disabled(tmp_path):
    client = Client(Runtime(tmp_path))
    client.open()
    summaries = client.call("module.list")["result"]
    assert summaries[0]["enabled"] is False and "input_schema" not in summaries[0]
    capabilities = client.call("capability.list", {"module_id": "office.xlsx"})["result"]
    assert len(capabilities) == 3 and all("input_schema" not in c for c in capabilities)
    detail = client.call("capability.get", {"capability_id": "xlsx.update_cells"})["result"]
    assert detail["input_schema"]["required"] and len(detail["contract_sha256"]) == 64
    assert detail["recovery"]["automatic_retry"] is False
    assert client.call("capability.get", {"capability_id": "unknown"})["status"] == "denied"


def test_real_update_read_and_independent_verification(setup):
    client, path, _ = setup
    from openpyxl import load_workbook
    before = path.read_bytes()
    output = path.parent / "updated.xlsx"
    req = request(client, "xlsx.update_cells", inputs(path, output_path=str(output), cells={"A2": 370, "B2": "新负责人"}))
    queued, final = invoke(client, req, "update")
    assert final["status"] == "completed", final
    result = final["result"]["output"]
    assert result["status"] == "verified" and result["evidence"]["reopened_cells_match"]
    assert path.read_bytes() == before
    book = load_workbook(output, read_only=True)
    assert book["Summary"]["A2"].value == 370 and book["Summary"]["B2"].value == "新负责人"
    book.close()
    _, read = invoke(client, request(client, "xlsx.read_range", inputs(output, range="A1:B2")))
    assert read["result"]["output"]["values"]["A2"] == 370
    _, verified = invoke(client, request(client, "xlsx.verify_cells", inputs(output, cells={"A2": 370})))
    assert verified["result"]["output"]["status"] == "verified"
    _, wrong = invoke(client, request(client, "xlsx.verify_cells", inputs(output, cells={"A2": 999})))
    assert wrong["status"] == "failed"
    # Repeat retrieves the terminal receipt and does not try to publish again.
    cached = client.call("capability.invoke", req, "update")
    assert cached["result"]["job_id"] == queued["result"]["job_id"]
    events = client.runtime.store.query(operation_id="update")
    assert len([e for e in events if e["kind"] == "job.started"]) == 1
    assert any(e["kind"] == "capability.resolved" for e in events)
    assert "新负责人" not in json.dumps(events, ensure_ascii=False)


@pytest.mark.parametrize("change", ["module", "trusted", "write", "version", "digest", "implementation", "paths"])
def test_grants_versions_and_implementations_refuse_before_job(setup, change):
    client, path, _ = setup
    req = request(client, "xlsx.read_range", inputs(path, range="A1:B2"))
    if change == "module":client.runtime.policy["modules"] = []
    if change == "trusted":client.runtime.policy["trusted_host"] = False
    if change == "write":client.runtime.policy["write"] = False
    if change == "paths":client.runtime.policy["paths"] = []
    if change == "version":req["version"] = "2.0.0"
    if change == "digest":req["contract_sha256"] = "0" * 64
    if change == "implementation":req["implementation"] = "arbitrary.module"
    result = client.call("capability.invoke", req)
    assert result["status"] == "denied" and result["dispatched"] is False
    assert not client.runtime.store.query(kind="job.started")


@pytest.mark.parametrize("version", [None, "2.0.0", "4.0.0", "3.1.5rc1"])
def test_dependency_health_and_stale_contract_are_refused(setup, monkeypatch, version):
    client, path, _ = setup
    req = request(client, "xlsx.read_range", inputs(path, range="A1"))
    def probe(name):
        if version is None:raise importlib.metadata.PackageNotFoundError(name)
        return version
    monkeypatch.setattr(registry.importlib.metadata, "version", probe)
    assert client.call("module.list")["result"][0]["health"]["available"] is False
    req = request(client, "xlsx.read_range", inputs(path, range="A1"))
    result = client.call("capability.invoke", req)
    assert result["status"] == "denied" and result["dispatched"] is False


def test_approval_exact_inputs_and_source_change_before_resume(setup):
    client, path, _ = setup
    client.runtime.policy["require_approval"] = True
    output = path.parent / "approved.xlsx"
    req = request(client, "xlsx.update_cells", inputs(path, output_path=str(output), cells={"A2": 370}))
    pending = client.call("capability.invoke", req, "approve")
    assert pending["status"] == "approval_required" and not output.exists()
    changed = copy.deepcopy(req);changed["inputs"]["cells"]["A2"] = 999
    assert client.call("capability.invoke", changed, "approve")["status"] == "denied"
    client.runtime.store.decide(pending["approval_id"], "approved", "owned fixture")
    invoke(client, req, "approve")
    assert output.exists()
    req = request(client, "xlsx.read_range", inputs(path, range="A1"))
    pending = client.call("capability.invoke", req, "changed-source")
    client.runtime.store.decide(pending["approval_id"], "approved", "owned fixture")
    path.write_bytes(path.read_bytes() + b"changed")
    refused = client.call("capability.invoke", req, "changed-source")
    assert refused["status"] == "denied" and refused["dispatched"] is False


@pytest.mark.parametrize("invalid", ["overwrite", "outside", "formula", "range", "boolean-coordinate", "unknown-input"])
def test_input_boundaries_zero_dispatch(setup, invalid):
    client, path, _ = setup
    cid = "xlsx.update_cells"
    data = inputs(path, output_path=str(path.parent / "output.xlsx"), cells={"A2": 370})
    if invalid == "overwrite":data["output_path"] = str(path)
    if invalid == "outside":data["output_path"] = str(path.parent.parent / "outside.xlsx")
    if invalid == "formula":data["cells"] = {"A2": "=SUM(A1:A2)"}
    if invalid == "boolean-coordinate":data["cells"] = {"XFE1": True}
    if invalid == "unknown-input":data["extra"] = "override"
    if invalid == "range":cid = "xlsx.read_range";data = inputs(path, range="A1:XFD1048576")
    result = client.call("capability.invoke", request(client, cid, data))
    assert result["status"] == "denied" and result["dispatched"] is False


def test_existing_output_and_formula_workbook_preserve_bytes(setup):
    client, path, _ = setup
    from openpyxl import load_workbook
    output = path.parent / "existing.xlsx";output.write_bytes(b"keep")
    req = request(client, "xlsx.update_cells", inputs(path, output_path=str(output), cells={"A2": 370}))
    assert client.call("capability.invoke", req)["status"] == "denied" and output.read_bytes() == b"keep"
    book = load_workbook(path);book["Summary"]["C2"] = "=A2+1";book.save(path);book.close()
    before = path.read_bytes();output = path.parent / "formula.xlsx"
    req = request(client, "xlsx.update_cells", inputs(path, output_path=str(output), cells={"A2": 370}))
    _, final = invoke(client, req)
    assert final["status"] == "failed" and not output.exists() and path.read_bytes() == before


def test_unknown_operation_never_launches_replacement(setup):
    client, path, _ = setup
    req = request(client, "xlsx.read_range", inputs(path, range="A1"))
    queued, final = invoke(client, req, "lost")
    with client.runtime.store.connect() as db:
        db.execute("UPDATE operations SET status='running', result=NULL WHERE id='lost'")
    result = client.call("capability.invoke", req, "lost")
    assert result["status"] == "completion_unknown"
    assert len(client.runtime.store.query(operation_id="lost", kind="job.started")) == 1


def test_module_not_added_to_existing_session_by_policy_expansion(setup):
    client, path, _ = setup
    with client.runtime.store.connect() as db:
        row = db.execute("SELECT policy FROM sessions WHERE id=?", (client.session_id,)).fetchone()
        scope = json.loads(row["policy"]);scope["modules"] = []
        db.execute("UPDATE sessions SET policy=? WHERE id=?", (json.dumps(scope), client.session_id))
    req = request(client, "xlsx.read_range", inputs(path, range="A1"))
    assert client.call("capability.invoke", req)["status"] == "denied"


def test_output_created_during_publish_is_never_overwritten(setup, monkeypatch):
    client, path, _ = setup
    from bokkio.modules import office_xlsx
    output = path.parent / "racing.xlsx"
    req = request(client, "xlsx.update_cells", inputs(path, output_path=str(output), cells={"A2": 370}))
    job = registry.prepare(client.runtime, client.runtime._session(client.session_id), req)
    monkeypatch.setattr(Client, "from_environment", lambda: client)
    original = office_xlsx.os.link
    def racing(source, destination):
        Path(destination).write_bytes(b"another writer")
        original(source, destination)
    monkeypatch.setattr(office_xlsx.os, "link", racing)
    with pytest.raises(FileExistsError):office_xlsx.execute(job["capability"], req["inputs"])
    assert output.read_bytes() == b"another writer"
    assert not list(path.parent.glob(".bokkio-xlsx-*"))


def test_post_publish_verification_failure_returns_unknown(setup, monkeypatch):
    client, path, _ = setup
    import openpyxl
    from bokkio.modules import office_xlsx
    output = path.parent / "published.xlsx"
    req = request(client, "xlsx.update_cells", inputs(path, output_path=str(output), cells={"A2": 370}))
    job = registry.prepare(client.runtime, client.runtime._session(client.session_id), req)
    monkeypatch.setattr(Client, "from_environment", lambda: client)
    original = openpyxl.load_workbook
    def broken(filename, **kwargs):
        if filename == output:raise OSError("Lost published-file observation")
        return original(filename, **kwargs)
    monkeypatch.setattr(openpyxl, "load_workbook", broken)
    result = office_xlsx.execute(job["capability"], req["inputs"])
    assert result["status"] == "completion_unknown" and result["evidence"]["published"]
    book = original(output);assert book["Summary"]["A2"].value == 370;book.close()


@pytest.mark.parametrize("stdout", ["invalid JSON", '{}', '{"status":"verified"}', '{"status":"completion_unknown"}'])
def test_missing_or_unknown_job_result_never_replays(setup, monkeypatch, stdout):
    client, path, _ = setup
    req = request(client, "xlsx.read_range", inputs(path, range="A1"))
    original = registry.prepare
    def bad_result(*args):
        job = original(*args);job["code"] = "print(" + repr(stdout) + ")";return job
    monkeypatch.setattr(registry, "prepare", bad_result)
    queued, final = invoke(client, req, "unknown-result")
    assert final["status"] == "completion_unknown"
    assert client.call("capability.invoke", req, "unknown-result")["status"] == "completion_unknown"
    assert len(client.runtime.store.query(operation_id="unknown-result", kind="job.started")) == 1


def test_mcp_cli_sdk_discovery_and_invocation_share_contract(setup, tmp_path):
    pytest.importorskip("mcp")
    client, path, policy = setup
    policy_path = tmp_path / "policy.json";policy_path.write_text(json.dumps(policy))
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    async def check():
        params = StdioServerParameters(command=sys.executable, args=["-m", "bokkio.runtime_cli", "--state", str(client.runtime.store.root), "--policy", str(policy_path), "serve"])
        async with stdio_client(params) as streams:
            async with ClientSession(*streams) as mcp:
                await mcp.initialize()
                tools = {t.name: t for t in (await mcp.list_tools()).tools}
                assert {"module.list", "capability.get", "capability.invoke"} <= set(tools)
                req = request(client, "xlsx.read_range", inputs(path, range="A1:B2"))
                queued = await mcp.call_tool("capability.invoke", {"arguments": req, "session_id": client.session_id, "operation_id": "mcp-module"})
                assert not queued.isError, queued
                final = client.wait_job(queued.structuredContent["result"]["job_id"], 20)
                assert final["result"]["result"]["output"]["values"]["A2"] == 320
                cli = subprocess.run([sys.executable, "-m", "bokkio.runtime_cli", "--state", str(client.runtime.store.root), "--policy", str(policy_path), "call", "capability.get", "--session", client.session_id, "--args", json.dumps({"capability_id": "xlsx.read_range"})], capture_output=True, text=True, timeout=15)
                assert cli.returncode == 0, cli.stderr
                assert json.loads(cli.stdout)["result"]["contract_sha256"] == req["contract_sha256"]
    asyncio.run(check())
