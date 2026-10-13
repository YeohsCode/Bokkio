"""Disposable XLSX comparison; no LLM, desktop, customer files or business scoring."""
import argparse
import hashlib
import json
from pathlib import Path
import platform
import time

from openpyxl import Workbook, load_workbook

from bokkio.runtime import Runtime
from bokkio.sdk import Client


def run(output):
    output.mkdir(parents=True, exist_ok=False)
    project = Path(__file__).resolve().parents[1]
    files = ["src/bokkio/"+name for name in ("runtime.py", "executor.py", "tool_schema.py", "model.py", "xa11y_backend.py", "capability_registry.py",
             "modules/office_xlsx.py", "modules/office_xlsx_contract.py")]
    files += ["scripts/verify_capability_modules.py", "pyproject.toml", "uv.lock"]
    source_hashes = {name: hashlib.sha256((project/name).read_bytes()).hexdigest() for name in files}
    policy = {"paths": [str(output.resolve())], "modules": ["office.xlsx"],
              "write": True, "trusted_host": True, "require_approval": False}
    client = Client(Runtime(output / "state", policy));client.open()
    rows = []
    try:
        assert client.call("module.list")["result"][0]["enabled"]
        assert len(client.call("capability.list")["result"]) == 3
        contract = client.call("capability.get", {"capability_id": "xlsx.update_cells"})["result"]
        for index in range(3):
            source = output / ("source-" + str(index) + ".xlsx")
            book = Workbook();book.active.title = "Summary";book.active.append(["Amount", "Owner"])
            book.active.append([320+index, "Fixture"]);book.save(source);book.close()
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            expected = {"A2": 370+index, "B2": "Updated fixture " + str(index)}
            trials = []
            for route in ("code", "module"):
                destination = output / (route + "-" + str(index) + ".xlsx")
                started = time.monotonic()
                if route == "code":
                    # The generic executor can already do this in one supplied script.
                    # This is a controlled baseline, not an autonomous Harness attempt.
                    code = "\n".join([
                        "from pathlib import Path", "from openpyxl import load_workbook", "import hashlib,os,tempfile",
                        "source=Path(" + repr(str(source.resolve())) + ")",
                        "output=Path(" + repr(str(destination.resolve())) + ")",
                        "expected=" + repr(expected), "before=" + repr(before),
                        "assert hashlib.sha256(source.read_bytes()).hexdigest()==before",
                        "book=load_workbook(source)",
                        "for cell,value in expected.items():book['Summary'][cell]=value",
                        "fd,tmp=tempfile.mkstemp(dir=output.parent,suffix='.xlsx');os.close(fd)",
                        "try:", "    book.save(tmp);os.link(tmp,output)", "finally:", "    Path(tmp).unlink(missing_ok=True);book.close()",
                        "check=load_workbook(output,read_only=True)",
                        "assert all(check['Summary'][cell].value==value for cell,value in expected.items())",
                        "check.close()", "assert hashlib.sha256(source.read_bytes()).hexdigest()==before", "print('verified')"])
                    queued = client.call("code.run", {"code": code, "cwd": str(output.resolve())}, "baseline-"+str(index))
                else:
                    queued = client.call("capability.invoke", {"capability_id": contract["id"], "version": contract["version"],
                        "contract_sha256": contract["contract_sha256"], "inputs": {"path": str(source.resolve()), "sheet": "Summary",
                        "expected_sha256": before, "output_path": str(destination.resolve()), "cells": expected}}, "module-"+str(index))
                assert queued["status"] == "queued", queued
                final = client.wait_job(queued["result"]["job_id"], 30)
                assert final["result"]["status"] == "completed", final
                if route == "module":assert final["result"]["result"]["output"]["status"] == "verified"
                # Independent artifact read; do not rely on worker exit or its own verification.
                independent = load_workbook(destination, read_only=True)
                matched = all(independent["Summary"][cell].value == value for cell,value in expected.items())
                independent.close()
                unchanged = hashlib.sha256(source.read_bytes()).hexdigest() == before
                assert matched and unchanged
                trials.append({"route": route, "passed": True, "elapsed_ms": round((time.monotonic()-started)*1000, 2),
                    "execution_requests": 1, "independent_cells_match": matched, "source_unchanged": unchanged,
                    "output_sha256": hashlib.sha256(destination.read_bytes()).hexdigest()})
            rows.append({"sample": index+1, "trials": trials})
        report = {"schema": "bokkio.capability-evaluation.v1", "platform": platform.system(), "samples": rows,
            "source_files": source_hashes,
            "passed": 6, "total": 6, "module_contract_sha256": contract["contract_sha256"],
            "scope": "controlled_file_operations_not_autonomous_business_or_GUI",
            "limitations": ["baseline includes supplied exact code", "no LLM tokens or business success measurement", "latency is descriptive, not a performance benchmark", "Windows not executed"]}
        assert source_hashes == {name: hashlib.sha256((project/name).read_bytes()).hexdigest() for name in files}, "Source changed during comparison"
        (output / "summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=2)+"\n")
        return report
    finally:
        client.call("session.close")


if __name__ == "__main__":
    parser = argparse.ArgumentParser();parser.add_argument("--output", type=Path, required=True)
    print(json.dumps(run(parser.parse_args().output), ensure_ascii=False))
